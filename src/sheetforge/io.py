"""Read and write tables as CSV, JSON or XLSX.

Encoding fallback for text inputs: utf-8-sig → utf-8 → latin-1.
"""

from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

ENCODINGS: tuple[str, ...] = ("utf-8-sig", "utf-8", "latin-1")


@dataclass(slots=True)
class Table:
    columns: list[str]
    rows: list[dict[str, Any]]

    def __len__(self) -> int:
        return len(self.rows)

    def column(self, name: str) -> str:
        """Resolve a column name case-insensitively."""
        for candidate in self.columns:
            if candidate.lower() == name.lower():
                return candidate
        raise ValueError(f"unknown column: {name!r} (available: {', '.join(self.columns)})")


def decode(data: bytes) -> tuple[str, str]:
    for encoding in ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError("unable to decode input file with utf-8-sig, utf-8 or latin-1")


def read_table(path: str | Path) -> Table:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"input file not found: {source}")
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(source)
    if suffix == ".json":
        return _read_json(source)
    return _read_delimited(source, delimiter="\t" if suffix == ".tsv" else ",")


def _read_delimited(path: Path, delimiter: str = ",") -> Table:
    text, _ = decode(path.read_bytes())
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        raise ValueError(f"empty input file: {path}") from None
    columns = [cell.strip() or f"col{i + 1}" for i, cell in enumerate(header)]
    rows: list[dict[str, Any]] = []
    for raw in reader:
        if not any(cell.strip() for cell in raw):
            continue
        rows.append({columns[i]: (raw[i] if i < len(raw) else "") for i in range(len(columns))})
    return Table(columns=columns, rows=rows)


def _read_json(path: Path) -> Table:
    text, _ = decode(path.read_bytes())
    data = json.loads(text)
    if not isinstance(data, list) or not data:
        raise ValueError(f"expected a non-empty JSON array of objects: {path}")
    columns: list[str] = []
    for item in data:
        if not isinstance(item, dict):
            raise ValueError(f"expected a JSON array of objects: {path}")
        for key in item:
            if key not in columns:
                columns.append(key)
    rows = [{col: item.get(col, "") for col in columns} for item in data]
    return Table(columns=columns, rows=rows)


def _read_xlsx(path: Path) -> Table:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            raise ValueError(f"empty input file: {path}")
        columns = [
            str(cell).strip() if cell is not None else f"col{i + 1}"
            for i, cell in enumerate(header)
        ]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if values is None or all(value is None for value in values):
                continue
            row: dict[str, Any] = {}
            for i, col in enumerate(columns):
                value = values[i] if i < len(values) else None
                row[col] = "" if value is None else value
            rows.append(row)
    finally:
        workbook.close()
    return Table(columns=columns, rows=rows)


def to_csv(table: Table, path: str | Path) -> Path:
    target = Path(path)
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=table.columns, extrasaction="ignore")
        writer.writeheader()
        for row in table.rows:
            writer.writerow({col: row.get(col, "") for col in table.columns})
    return target


def to_json(table: Table, path: str | Path) -> Path:
    target = Path(path)
    payload = [{col: row.get(col, "") for col in table.columns} for row in table.rows]
    target.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return target


def to_xlsx(table: Table, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(table.columns)
    for row in table.rows:
        sheet.append([row.get(col, "") for col in table.columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    workbook.save(target)
    return target


def write_table(table: Table, path: str | Path) -> Path:
    """Dispatch on file suffix (.csv / .json / .xlsx)."""
    target = Path(path)
    suffix = target.suffix.lower()
    if suffix == ".csv":
        return to_csv(table, target)
    if suffix == ".json":
        return to_json(table, target)
    if suffix in {".xlsx", ".xlsm"}:
        return to_xlsx(table, target)
    raise ValueError(f"unsupported format: {suffix!r} (use .csv, .json or .xlsx)")


def write_rows(rows: list[dict[str, Any]], columns: list[str], path: str | Path) -> Path:
    return to_csv(Table(columns=columns, rows=rows), path)
